"""简历文件解析：PDF / DOCX → 纯文本全文。

三级策略（全离线，模型随 runtime 打包）：
    DOCX  → python-docx（段落 + 表格）
    PDF   → PyMuPDF 文本层（电子版简历秒出）
    PDF   → 文本层为空（扫描件 / 图片版）自动切换 RapidOCR
            （onnxruntime 引擎，页面渲染后逐页识别，按行合并）
"""
from __future__ import annotations

import re
import time
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class ExtractResult:
    """一次文档解析的结果。"""

    text: str = ""
    engine: str = ""            # docx / pdf_text / pdf_ocr
    pages: int = 0
    ocr_seconds: float = 0.0
    warnings: list[str] = field(default_factory=list)


# 文本层少于该字符数视为「没有文本层」，走 OCR
_TEXT_LAYER_MIN_CHARS = 40

_ocr_engine = None      # 懒加载的 RapidOCR 单例（模型加载约 1-2 秒）


def _get_ocr():
    """懒加载 RapidOCR 引擎（离线模型，随 runtime 分发）。"""
    global _ocr_engine
    if _ocr_engine is None:
        from rapidocr_onnxruntime import RapidOCR

        _ocr_engine = RapidOCR()
    return _ocr_engine


def _normalize(text: str) -> str:
    """统一换行、去 BOM、压缩行尾空白。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\ufeff", "")
    lines = [ln.rstrip() for ln in text.split("\n")]
    # 连续 3+ 空行压缩为 2 个
    out: list[str] = []
    blank = 0
    for ln in lines:
        if ln.strip():
            blank = 0
            out.append(ln)
        else:
            blank += 1
            if blank <= 2:
                out.append("")
    return "\n".join(out).strip()


# ---------------------------------------------------------------- DOCX

def _extract_docx(path: Path) -> ExtractResult:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = "  ".join([c for c in cells if c])
            if line:
                parts.append(line)
    return ExtractResult(
        text=_normalize("\n".join(parts)),
        engine="docx",
        pages=1,
    )


# ---------------------------------------------------------------- PDF 文本层

def _extract_pdf_text(path: Path) -> ExtractResult:
    import fitz  # PyMuPDF

    with fitz.open(str(path)) as doc:
        pages_text = [page.get_text("text") or "" for page in doc]
        pages = len(doc)
    joined = "\n".join(pages_text)
    return ExtractResult(
        text=_normalize(joined),
        engine="pdf_text",
        pages=pages,
    )


# ---------------------------------------------------------------- PDF OCR

def _merge_ocr_lines(raw: list) -> str:
    """把 OCR 的行框按阅读顺序合并成文本行。

    raw: [[box(4点), text, score], ...]
    同一视觉行（y 接近）内的多个框按 x 排序拼接，宽间距补两个空格。
    """
    items = []
    for box, text, _score in raw:
        text = (text or "").strip()
        if not text:
            continue
        ys = [pt[1] for pt in box]
        xs = [pt[0] for pt in box]
        items.append((sum(ys) / len(ys), min(xs), max(xs), text))
    items.sort(key=lambda t: (t[0], t[1]))

    lines: list[list[tuple[float, float, float, str]]] = []
    for item in items:
        if lines and abs(item[0] - lines[-1][0][0]) <= 14:
            lines[-1].append(item)
        else:
            lines.append([item])

    out: list[str] = []
    for line in lines:
        line.sort(key=lambda t: t[1])
        buf = line[0][3]
        prev_max_x = line[0][2]
        for item in line[1:]:
            gap = item[1] - prev_max_x
            sep = "  " if gap >= 30 else (" " if gap >= 8 else "")
            buf += sep + item[3]
            prev_max_x = max(prev_max_x, item[2])
        out.append(buf)
    return "\n".join(out)


def _extract_pdf_ocr(path: Path, dpi: int = 200) -> ExtractResult:
    import fitz  # PyMuPDF

    engine = _get_ocr()
    t0 = time.time()
    all_lines: list[str] = []
    warnings: list[str] = []
    with fitz.open(str(path)) as doc:
        pages = len(doc)
        for i, page in enumerate(doc):
            pix = page.get_pixmap(dpi=dpi)
            img_bytes = pix.tobytes("png")
            result, _elapse = engine(img_bytes)
            if not result:
                warnings.append(f"第 {i + 1} 页未识别到文字")
                continue
            if i > 0:
                all_lines.append("")     # 页间空行
            all_lines.append(_merge_ocr_lines(result))
    ocr_seconds = time.time() - t0
    return ExtractResult(
        text=_normalize("\n".join(all_lines)),
        engine="pdf_ocr",
        pages=pages,
        ocr_seconds=round(ocr_seconds, 1),
        warnings=warnings,
    )


# ---------------------------------------------------------------- 对外入口

def extract_text(path: str | Path) -> ExtractResult:
    """解析简历文件，返回全文（自动选择文本层 / OCR）。"""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".docx":
        return _extract_docx(p)
    if suffix == ".pdf":
        result = _extract_pdf_text(p)
        if len(result.text) >= _TEXT_LAYER_MIN_CHARS:
            return result
        # 图片版 / 扫描件 → OCR
        ocr = _extract_pdf_ocr(p)
        ocr.pages = result.pages
        if not ocr.text:
            ocr.warnings.append("OCR 未识别出任何文字，请确认文件清晰度")
        return ocr
    raise ValueError(f"不支持的文件类型：{suffix}（仅支持 .pdf / .docx）")
